"""Convert every source file (PDF, PPTX, DOCX, XLSX) to a Markdown text twin.

- Course files: `<name>.<ext>` -> `<name>.md` next to the original, with page/slide markers.
- Textbooks: split per chapter into `Textbooks/<book> (md)/NN - <title>.md`,
  with `<!-- pdf p. N -->` markers so agents can cite pages.

Originals are never modified. Re-run safely: existing .md twins are overwritten.
Usage: python tools/convert_to_md.py   (from the repo root)
"""
import re
from pathlib import Path

import fitz  # PyMuPDF
import docx
import openpyxl
from pptx import Presentation

ROOT = Path(__file__).resolve().parent.parent
SKIP_DIRS = {".git", "tools", "map", "progress"}


def clean(t: str) -> str:
    t = t.replace(" ", " ")
    t = re.sub(r"[ \t]+\n", "\n", t)
    return re.sub(r"\n{3,}", "\n\n", t).strip()


def pdf_to_md(p: Path) -> str:
    d = fitz.open(p)
    out = [f"# {p.stem}\n", f"_Source: `{p.name}` ({d.page_count} pages). Text extraction; formulas/figures may be imperfect - check the PDF._\n"]
    for i, page in enumerate(d, 1):
        out.append(f"\n<!-- page {i} -->\n## Page {i}\n\n{clean(page.get_text())}\n")
    return "\n".join(out)


def pptx_to_md(p: Path) -> str:
    prs = Presentation(p)
    out = [f"# {p.stem}\n", f"_Source: `{p.name}` ({len(prs.slides)} slides). Text only; check the deck for figures._\n"]
    for i, s in enumerate(prs.slides, 1):
        texts = []
        for sh in s.shapes:
            if sh.has_text_frame:
                texts.append(sh.text_frame.text)
            elif getattr(sh, "has_table", False) and sh.has_table:
                for r in sh.table.rows:
                    texts.append(" | ".join(c.text for c in r.cells))
        notes = s.notes_slide.notes_text_frame.text if s.has_notes_slide else ""
        body = clean("\n".join(texts))
        out.append(f"\n<!-- slide {i} -->\n## Slide {i}\n\n{body}\n")
        if notes.strip():
            out.append(f"\n> Notes: {clean(notes)}\n")
    return "\n".join(out)


def docx_to_md(p: Path) -> str:
    d = docx.Document(p)
    out = [f"# {p.stem}\n", f"_Source: `{p.name}`._\n"]
    for para in d.paragraphs:
        t = para.text.strip()
        if not t:
            continue
        style = ((para.style.name if para.style is not None else "") or "").lower()
        if style.startswith("heading"):
            lvl = re.findall(r"\d", style)
            out.append("#" * (int(lvl[0]) + 1 if lvl else 2) + " " + t)
        else:
            out.append(t)
    for ti, tbl in enumerate(d.tables, 1):
        out.append(f"\n**Table {ti}**\n")
        for r in tbl.rows:
            out.append("| " + " | ".join(c.text.strip() for c in r.cells) + " |")
    return clean("\n\n".join(out))


def xlsx_to_md(p: Path) -> str:
    wb = openpyxl.load_workbook(p, data_only=True)
    out = [f"# {p.stem}\n", f"_Source: `{p.name}`._\n"]
    for ws in wb:
        out.append(f"\n## Sheet: {ws.title}\n")
        for r in ws.iter_rows(values_only=True):
            if any(c is not None for c in r):
                out.append("| " + " | ".join("" if c is None else str(c) for c in r) + " |")
    return "\n".join(out)


def split_book(p: Path, outdir: Path, min_level: int):
    d = fitz.open(p)
    toc = [(l, t, pg) for l, t, pg in d.get_toc() if l == min_level and re.match(r"^\d+[\s:]", t)]
    outdir.mkdir(exist_ok=True)
    ends = [pg for _, _, pg in toc[1:]]
    # last chapter ends at next top-level TOC entry after it
    last_start = toc[-1][2]
    later = sorted(pg for l, _, pg in d.get_toc() if l <= min_level and pg > last_start)
    ends.append(later[0] if later else d.page_count + 1)
    index = [f"# {p.stem} - chapters\n", "| Ch | Title | PDF pages | File |", "|---|---|---|---|"]
    for (lvl, title, start), end in zip(toc, ends):
        num = int(re.match(r"^(\d+)", title).group(1))
        name = re.sub(r"[\x00-\x1f]", "", re.sub(r"^\d+[\s:]+", "", title)).replace("�", "'").strip()
        safe = re.sub(r'[\\/:*?"<>|]', "-", name)[:60].strip()
        fn = f"{num:02d} - {safe}.md"
        parts = [f"# Chapter {num}: {name}\n", f"_Source: `{p.name}`, PDF pages {start}-{end - 1}. Page markers below are PDF page numbers._\n"]
        for i in range(start - 1, end - 1):
            parts.append(f"\n<!-- pdf p. {i + 1} -->\n{clean(d[i].get_text())}\n")
        (outdir / fn).write_text("\n".join(parts), encoding="utf-8")
        index.append(f"| {num} | {name} | {start}-{end - 1} | [{fn}](<{fn}>) |")
    (outdir / "_index.md").write_text("\n".join(index) + "\n", encoding="utf-8")
    return len(toc)


def main():
    conv = {".pdf": pdf_to_md, ".pptx": pptx_to_md, ".docx": docx_to_md, ".xlsx": xlsx_to_md}
    n = 0
    for f in sorted(ROOT.rglob("*")):
        if not f.is_file() or set(f.relative_to(ROOT).parts) & SKIP_DIRS:
            continue
        if f.parent.name == "Textbooks" or f.suffix.lower() not in conv:
            continue
        f.with_suffix(".md").write_text(conv[f.suffix.lower()](f), encoding="utf-8")
        n += 1
    tb = ROOT / "Textbooks"
    a = split_book(tb / "Rosen & Gayer - Public Finance (10th ed).pdf", tb / "Rosen & Gayer (md)", 1)
    b = split_book(tb / "Personnel Economics.pdf", tb / "Kuhn - Personnel Economics (md)", 2)
    print(f"converted {n} course files; Rosen&Gayer {a} chapters; Kuhn {b} chapters")


if __name__ == "__main__":
    main()
